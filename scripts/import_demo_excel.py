"""Import the three-sheet collection workbook without changing flow or logo assets.

Run with the project's documented Python environment. Requires openpyxl.
The importer writes a report even on failure and never writes output with errors.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import date, datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import sys
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_BASE = ROOT / "public/assets/demo/demo-data.json"
MISSING = {"", "缺", "—", "--", "-", "暂无", "待补充", "待接入", "待核实", "未找到", "n/a", "na", "null"}
CHAIN_ALIASES = {
    "chemicals": ["化工新材料（含盐化工）", "化工新材料", "盐化工", "化工", "新材料化工"],
    "new-energy": ["新能源及材料", "新能源", "新能源材料", "锂电新材料", "锂电池材料"],
    "electronics": ["电子信息与半导体", "电子信息", "半导体", "光电子信息"],
    "automotive": ["汽车及零部件", "汽车零部件", "汽车制造", "汽车"],
    "equipment": ["高端装备", "装备制造", "高端装备制造", "智能装备"],
    "biomedicine": ["生物医药", "医药", "生命健康"],
    "food": ["食品及农产品加工", "食品加工", "农产品加工", "食品"],
    "textile": ["纺织服装", "纺织", "服装"],
    "digital": ["数字经济", "软件信息服务", "数字产业"],
    "future": ["未来产业（低空经济、人工智能等）", "未来产业", "低空经济", "人工智能"],
}


def norm(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).replace("（", "(").replace("）", ")")


def plain(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def missing(value: Any) -> bool:
    return value is None or (isinstance(value, str) and (value.strip().lower() in MISSING or re.match(r"^(缺|待补充|待接入|待核实|未找到|暂无)[（(：:]", value.strip()) is not None))


def split_list(value: Any) -> list[str]:
    """Split enumeration separators outside parentheses, keeping original labels."""
    if missing(value):
        return []
    result, current, depth = [], [], 0
    for char in plain(value):
        if char in "（(": depth += 1
        elif char in "）)": depth = max(0, depth - 1)
        if depth == 0 and char in "、,，;；\n":
            if "".join(current).strip(): result.append("".join(current).strip())
            current = []
        else:
            current.append(char)
    if "".join(current).strip(): result.append("".join(current).strip())
    return list(dict.fromkeys(result))


def stable_id(prefix: str, value: str) -> str:
    return prefix + "-" + hashlib.sha1(value.encode("utf-8")).hexdigest()[:12]


class Importer:
    def __init__(self, path: Path, base: dict[str, Any], year: int = 2025):
        self.path = path
        self.data = deepcopy(base)
        self.year = year
        self.book = load_workbook(path, data_only=False)
        self.cached = load_workbook(path, data_only=True)
        self.errors: list[dict[str, Any]] = []
        self.warnings: list[dict[str, Any]] = []
        self.counts = {"regions": 0, "chainStages": 0, "indicators": 0, "pillars": 0, "chainLinks": 0, "leaders": 0, "policies": 0, "newCompanies": 0}
        self.chain_lookup = {norm(name): cid for cid, names in CHAIN_ALIASES.items() for name in names}
        for chain in self.data["chains"]:
            self.chain_lookup[norm(chain["name"])] = chain["id"]
            self.chain_lookup[norm(chain["shortName"])] = chain["id"]
        self.parsed_chain_stages: set[tuple[str, str]] = set()

    def issue(self, cell: Any, message: str, *, error: bool = True):
        entry = {"sheet": getattr(getattr(cell, "parent", None), "title", ""), "cell": getattr(cell, "coordinate", ""), "message": message, "value": plain(getattr(cell, "value", None))}
        (self.errors if error else self.warnings).append(entry)

    def value(self, cell: Any) -> Any:
        if cell is None: return None
        if cell.data_type == "f":
            result = self.cached[cell.parent.title][cell.coordinate].value
            if result is None: self.issue(cell, "公式没有缓存结果；请在 Excel 中重新计算并保存，或粘贴为值")
            return result
        return cell.value

    def text(self, cell: Any, *, source: bool = False) -> str:
        value = self.value(cell)
        if missing(value): return ""
        text = plain(value)
        target = getattr(getattr(cell, "hyperlink", None), "target", None)
        if source and target and target not in text: text += " | " + target
        return text

    def number(self, cell: Any, *, percentage: bool = False, minimum: float | None = None, maximum: float | None = None, integer: bool = False) -> float | int | None:
        value = self.value(cell)
        if missing(value): return None
        if isinstance(value, bool):
            self.issue(cell, "布尔值不能作为数字"); return None
        percent_text = isinstance(value, str) and value.strip().endswith(("%", "％"))
        if isinstance(value, str):
            token = value.strip().replace(",", "").replace("，", "").replace("−", "-")
            if percent_text and percentage: token = token[:-1].strip()
            try: value = float(token)
            except ValueError:
                self.issue(cell, "数字格式错误。请填数字；GDP 单位为亿元，比例单位为百分比，不要混入“约”“万亿”等文字")
                return None
        if not isinstance(value, (int, float)) or not math.isfinite(value):
            self.issue(cell, "需要有限数值"); return None
        if percentage and not percent_text and "%" in getattr(cell, "number_format", ""):
            value *= 100
        if minimum is not None and value < minimum or maximum is not None and value > maximum:
            self.issue(cell, f"数字超出范围 {minimum if minimum is not None else '-∞'} 至 {maximum if maximum is not None else '∞'}")
            return None
        if integer and value != int(value):
            self.issue(cell, "该字段需要整数"); return None
        return int(value) if integer else round(float(value), 10)

    def sheet(self, name: str, alternatives: list[str] | None = None):
        names = [name] + (alternatives or [])
        for existing in self.book.sheetnames:
            if norm(existing) in [norm(n) for n in names]: return self.book[existing]
        self.errors.append({"sheet": name, "cell": "", "message": "缺少工作表", "value": ""})
        return None

    @staticmethod
    def fields(row: tuple[Any, ...]) -> dict[str, Any]:
        return {norm(cell.value): cell.column - 1 for cell in row if cell.value is not None}

    @staticmethod
    def field(row: tuple[Any, ...], header: dict[str, Any], *names: str):
        for name in names:
            index = header.get(norm(name))
            if index is not None and index < len(row): return row[index]
        return None

    def table(self, sheet: Any, required: list[tuple[str, ...]]):
        if sheet is None: return
        header = None
        for row in sheet.iter_rows():
            candidate = self.fields(row)
            if all(any(norm(name) in candidate for name in group) for group in required):
                header = candidate
                continue
            if header is not None and any(not missing(c.value) for c in row): yield row, header
        if header is None:
            self.errors.append({"sheet": sheet.title, "cell": "", "message": "没有找到表头：" + "、".join(x[0] for x in required), "value": ""})

    def chain(self, cell: Any, *, optional: bool = False) -> str | None:
        name = self.text(cell)
        if not name: return None
        cid = self.chain_lookup.get(norm(name))
        if not cid: self.issue(cell, f"未识别产业链“{name}”，请使用约定的 10 条统一名称", error=not optional)
        return cid

    def source(self, cell: Any, has_data: bool = True) -> str:
        source = self.text(cell, source=True)
        if has_data and not source: self.issue(cell, "这条记录有内容但没有来源，请补文件名或链接", error=False)
        return source

    def stage(self, chain_id: str, name: str, products: list[str] | None = None) -> dict[str, Any] | None:
        chain = next(x for x in self.data["chains"] if x["id"] == chain_id)
        # Name/ID wins over shared inputs such as lithium salts in two stages.
        for stage in chain["stages"]:
            labels = {norm(stage["id"]), norm(stage["name"])} | {norm(x) for x in stage["aliases"]}
            if norm(name) in labels: return stage
        keys = {norm(name)} | {norm(x) for x in (products or [])}
        scores = [(len(keys & {norm(x) for x in stage["products"]}), stage) for stage in chain["stages"]]
        maximum = max((score for score, _ in scores), default=0)
        best = [stage for score, stage in scores if score == maximum and score > 0]
        if len(best) == 1: return best[0]
        return None

    def import_chains(self):
        sheet = self.sheet("统一产业链名单", ["产业链名单"])
        for row, header in self.table(sheet, [("产业链",), ("环节", "环节（上/中/下游）"), ("包含哪些产品或业务", "产品或业务")]):
            get = lambda *names: self.field(row, header, *names)
            cid = self.chain(get("产业链"))
            if not cid:
                if not self.text(get("产业链")) and self.text(get("包含哪些产品或业务", "产品或业务")):
                    self.issue(get("产业链"), "填写了产品业务，但产业链名称为空")
                continue
            stage_cell = get("环节", "环节（上/中/下游）")
            position = self.text(stage_cell).replace("上/中/下游", "")
            position = {"上": "上游", "中": "中游", "下": "下游"}.get(position, position)
            if position not in ("上游", "中游", "下游"):
                self.issue(stage_cell, "环节位置应为上游、中游或下游"); continue
            products = split_list(self.value(get("包含哪些产品或业务", "产品或业务")))
            aliases = split_list(self.value(get("各地纲要里的常见说法", "常见说法")))
            name = self.text(get("环节名称")) or (products[0] if products else "")
            if not name:
                self.issue(get("包含哪些产品或业务", "产品或业务"), "产业环节缺少产品业务或环节名称"); continue
            existing = self.stage(cid, name, products)
            sid = existing["id"] if existing else stable_id(cid, position + name)
            if (cid, sid) in self.parsed_chain_stages:
                self.issue(get("产业链"), f"重复产业环节：{name}；请合并为一行"); continue
            self.parsed_chain_stages.add((cid, sid))
            update = dict(id=sid, stage=position, name=name, products=products, aliases=aliases)
            chain = next(x for x in self.data["chains"] if x["id"] == cid)
            if existing: existing.update(update)
            else: chain["stages"].append(update)
            for alias in aliases: self.chain_lookup.setdefault(norm(alias), cid)
            self.counts["chainStages"] += 1

    def import_regions(self):
        sheet = self.sheet("全国省市数据表", ["全国省市数据", "全国省市"])
        seen = set()
        for row, header in self.table(sheet, [("地区",), ("级别（省/市）", "级别"), ("所属省",), ("2025 年 GDP（亿元）", "2025年GDP（亿元）", "GDP（亿元）")]):
            get = lambda *names: self.field(row, header, *names)
            name = self.text(get("地区"))
            if not name: continue
            level = self.text(get("级别（省/市）", "级别"))
            if level not in ("省", "市"):
                self.issue(get("级别（省/市）", "级别"), "级别应为省或市"); continue
            def area_name(n: str): return norm(n).replace("省", "").replace("市", "").replace("壮族自治区", "").replace("回族自治区", "").replace("维吾尔自治区", "").replace("自治区", "")
            parent_name = self.text(get("所属省"))
            parent = next((x for x in self.data["regions"] if x["level"] == "省" and area_name(x["name"]) == area_name(parent_name)), None)
            if level == "市" and not parent:
                self.issue(get("所属省"), "城市需要填写可识别的所属省"); continue
            existing = next((x for x in self.data["regions"] if x["level"] == level and area_name(x["name"]) == area_name(name)), None)
            original = next((x for x in getattr(self, "original", self.data)["regions"] if x["level"] == level and area_name(x["name"]) == area_name(name)), None)
            lookup = existing or original
            if level == "省" and not lookup:
                self.issue(get("地区"), "省级地区不在约定的 31 个省级目录中"); continue
            rid = lookup["id"] if lookup else stable_id("city", (parent or {}).get("id", "") + name)
            if rid in seen:
                self.issue(get("地区"), "同一地区重复，请保留一行"); continue
            seen.add(rid)
            gdp = self.number(get("2025 年 GDP（亿元）", "2025年GDP（亿元）", "GDP（亿元）"), minimum=0)
            growth = self.number(get("GDP 增速（%）", "GDP增速（%）", "增速（%）"), percentage=True)
            shares = [self.number(get(k, k + "（%）"), percentage=True, minimum=0, maximum=100) for k in ("一产占比", "二产占比", "三产占比")]
            if all(v is not None for v in shares) and abs(sum(shares) - 100) > 1:
                self.issue(get("一产占比", "一产占比（%）"), "三产占比合计偏离 100%，请检查口径或输入", error=False)
            pillars = split_list(self.value(get("支柱产业", "支柱产业（原文）")))
            priorities = split_list(self.value(get("十五五重点产业", "十五五重点产业（原文）")))
            mapped = self.text(get("对应产业链"))
            chain_ids = []
            for token in split_list(mapped) if mapped else pillars + priorities:
                cid = self.chain_lookup.get(norm(token))
                if cid: chain_ids.append(cid)
                else: self.issue(get("对应产业链") if mapped else get("支柱产业"), f"未映射产业说法“{token}”；原文保留，请在对应产业链列补充统一名称", error=bool(mapped))
            has_data = any(v is not None for v in [gdp, growth] + shares) or bool(pillars or priorities)
            source = self.source(get("来源"), has_data)
            update = dict(id=rid, name=name, level=level, mapId=(lookup or {}).get("mapId", parent["id"] if parent else rid), year=self.year, gdp=gdp, growth=growth, shares=shares, pillars=pillars, priorities=priorities, chainIds=list(dict.fromkeys(chain_ids)), source=source, status="verified" if has_data else "missing")
            if level == "市": update["parentId"] = parent["id"]
            if existing: existing.update(update)
            else: self.data["regions"].append(update)
            self.counts["regions"] += 1

    def company_ids(self, cell: Any, chain_id: str | None, source: str) -> list[str]:
        ids = []
        for name in split_list(self.value(cell)):
            company = next((c for c in self.data["companies"] if norm(name) in (norm(c["name"]), norm(c["fullName"]))), None)
            if company is None:
                company = dict(id=stable_id("company", name), name=name, fullName=name, cityId="yunmeng", cityName="云梦县", logo=None, sector="待补充", summary="由云梦画像表接入；详细信息、标识与坐标待补充。", chainIds=[chain_id] if chain_id else [], stageIds=[], products=[], local=True, relationship="existing", shareholders=[], sources=[dict(label=source or "云梦画像 Excel")], status="missing")
                self.data["companies"].append(company); self.counts["newCompanies"] += 1
            ids.append(company["id"])
        return ids

    def import_profile(self):
        sheet = self.sheet("云梦画像")
        if sheet is None: return
        block = None; header = None; found = set(); seen: dict[str, set[Any]] = {}
        for row in sheet.iter_rows():
            candidate = self.fields(row)
            has = lambda *names: any(norm(n) in candidate for n in names)
            kind = None
            if has("年份", "年度") and has("GDP（亿元）", "GDP"): kind = "indicators"
            elif has("支柱产业") and has("产值规模", "产值规模（亿元）"): kind = "pillars"
            elif has("产业链") and has("强链/补链", "强链／补链", "状态"): kind = "chainLinks"
            elif has("产业链") and has("链长（姓名和职务）", "链长姓名", "链长"): kind = "leaders"
            elif has("文件名") and has("发布单位"): kind = "policies"
            if kind:
                block, header = kind, candidate; found.add(kind); seen.setdefault(kind, set()); continue
            if not block or not header or not any(not missing(c.value) for c in row): continue
            # Section labels separate blocks, so notes are never mistaken for rows.
            filled = [c for c in row if not missing(c.value)]
            if len(filled) == 1 and any(word in self.text(filled[0]) for word in ("指标", "支柱产业", "强链", "链长制", "政策文件")):
                block = None; header = None; continue
            get = lambda *names: self.field(row, header, *names)
            source = self.text(get("来源"), source=True)
            item = None; key = None
            if block == "indicators":
                year_cell = get("年份", "年度")
                if missing(self.value(year_cell)): continue
                year = self.number(year_cell, integer=True, minimum=1900, maximum=2200)
                if year is None: continue
                gdp = self.number(get("GDP（亿元）", "GDP"), minimum=0)
                growth = self.number(get("GDP增速（%）", "GDP 增速（%）", "增速（%）", "增速"), percentage=True)
                shares = [self.number(get(k, k + "（%）"), percentage=True, minimum=0, maximum=100) for k in ("一产占比", "二产占比", "三产占比")]
                count = self.number(get("规上工业企业数", "规上工业企业数（家）"), minimum=0, integer=True)
                populated = any(v is not None for v in [gdp, growth, count] + shares)
                source = self.source(get("来源"), populated)
                key = year
                item = dict(year=year, gdp=gdp, growth=growth, shares=shares, industrialEnterprises=count, source=source, status="verified" if populated else "missing")
            elif block == "pillars":
                name = self.text(get("支柱产业"))
                if not name: continue
                chain_cell = get("对应产业链", "产业链")
                if not self.text(chain_cell): chain_cell = get("支柱产业")
                cid = self.chain(chain_cell)
                if not cid: continue
                source = self.source(get("来源"))
                key = cid
                old = next((p for p in self.data["profile"]["pillars"] if p["chainId"] == cid), None)
                item = dict(id=old["id"] if old else stable_id("pillar", cid), name=name, chainId=cid, output=self.text(get("产值规模", "产值规模（亿元）")) or None, companyIds=self.company_ids(get("龙头企业", "企业名"), cid, source), park=self.text(get("所在园区")) or None, source=source)
            elif block == "chainLinks":
                cid = self.chain(get("产业链"))
                if not cid:
                    if not self.text(get("产业链")) and self.text(get("环节名称", "环节", "产品或业务")):
                        self.issue(get("产业链"), "填写了环节，但产业链名称为空")
                    continue
                stage_cell = get("环节名称", "环节", "产品或业务")
                stage = self.stage(cid, self.text(stage_cell))
                if not stage:
                    self.issue(stage_cell, "未识别产业环节，请与表二的环节名称或产品业务对应"); continue
                state_cell = get("强链/补链", "强链／补链", "状态")
                state = self.text(state_cell)
                state = {"强链": "present", "补链": "gap", "待核实": "unknown", "": "unknown"}.get(state, state)
                if state not in ("present", "gap", "unknown"):
                    self.issue(state_cell, "状态应为强链、补链或待核实"); continue
                source = self.source(get("来源"), state != "unknown")
                ids = self.company_ids(get("企业名", "企业名称", "企业"), cid, source)
                if state == "gap" and ids: self.issue(state_cell, "补链行同时填写了已有企业；请核实状态与环节覆盖范围")
                if state == "present" and not ids: self.issue(state_cell, "强链环节未填写企业名", error=False)
                key = stage["id"]
                pending = state == "gap" and bool(re.search(r"未找到|未核实|待核|缺少.*名录", source))
                item = dict(stageId=key, status="unknown" if pending else state, companyIds=ids, source=source, basis=plain(self.value(get("企业名", "企业名称", "企业"))))
                if state in ("present", "gap"):
                    item["reportedStatus"] = "强链" if state == "present" else "补链"
                    item["evidenceStatus"] = "pending" if pending else "reported"
            elif block == "leaders":
                cid = self.chain(get("产业链"))
                if not cid: continue
                person = self.text(get("链长姓名", "链长", "链长（姓名和职务）")) or None
                title = self.text(get("职务", "链长职务")) or None
                company = self.text(get("链主企业")) or None
                source = self.source(get("来源"), bool(person or title or company))
                key = cid; item = dict(chainId=cid, name=person, title=title, company=company, source=source)
            elif block == "policies":
                name = self.text(get("文件名"))
                if not name: continue
                link = self.text(get("链接"), source=True)
                if " | " in link: link = link.rsplit(" | ", 1)[-1]
                if link and not re.match(r"^https?://", link, re.I):
                    self.issue(get("链接"), "政策链接需要 http(s) 地址"); continue
                source = source or link or name
                key = name; item = dict(id=stable_id("policy", name), name=name, publisher=self.text(get("发布单位")), date=self.text(get("发布日期")), url=link or None, source=source)
            if item is None: continue
            if key in seen[block]:
                self.issue(row[0], "同一条画像记录重复，请合并为一行"); continue
            seen[block].add(key)
            items = self.data["profile"][block]
            key_field = {"indicators": "year", "pillars": "chainId", "chainLinks": "stageId", "leaders": "chainId", "policies": "name"}[block]
            old = next((x for x in items if x[key_field] == key), None)
            if old:
                if block == "chainLinks":
                    for field_name in ("reportedStatus", "evidenceStatus", "basis"): old.pop(field_name, None)
                old.update(item)
            else: items.append(item)
            self.counts[block] += 1
        for kind in ("indicators", "pillars", "chainLinks", "leaders", "policies"):
            if kind not in found: self.warnings.append({"sheet": sheet.title, "cell": "", "message": f"未找到画像分块 {kind}，保留原有内容", "value": ""})
        self.data["profile"]["indicators"].sort(key=lambda x: x["year"])

    def run(self) -> tuple[dict[str, Any], dict[str, Any]]:
        self.import_chains(); self.import_regions(); self.import_profile()
        self.data["meta"]["updatedAt"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
        self.data["meta"]["year"] = self.year
        self.data["meta"]["description"] = f"已从 {self.path.name} 接入采集字段；status=verified 表示接入用户提供来源，不代表独立核准。未填字段为 null。未导入的地区及原有城市匹配分数、企业点位、策略仍可能为示例，详见各记录来源。原始流空间资源保持独立。"
        report = {"input": str(self.path.resolve()), "counts": self.counts, "errors": self.errors, "warnings": self.warnings, "outputWritten": False}
        self.book.close(); self.cached.close()
        return self.data, report


class CollectedWorkbookImporter(Importer):
    """Full replacement for the delivered collection layout, retaining raw evidence.

    The workbook has merged chain cells and sector-level county evidence rather
    than the granular demonstration stages. No old demonstration gaps survive.
    """
    POSITION_IDS = {"上游": "upstream", "中游": "midstream", "下游": "downstream"}
    COMPANY_NAMES = {
        "gehua-huaxiang": ["湖北葛化华祥化学", "葛化华祥"],
        "lingyong-electronics": ["菱永电子"], "huasheng-xianghe": ["华盛锂电"],
        "tongyuan-materials": ["同源新材料"], "meilin-pharma": ["美林药业"],
        "tiangong-pharma": ["天宫制药"], "huaxin-silicone": ["华欣有机硅"],
        "annaiji": ["安耐吉"], "zhongyi": ["中一科技"],
        "tenglong-auto": ["腾龙汽配"], "kaite-electronics": ["开特电子"],
        "hangpeng-chemical": ["湖北航鹏化学动力"], "mixinyu-food": ["米芯语食品"],
        "yunxiangyuan": ["云香缘"], "yihao-food": ["壹号食品"], "yuntai-feed": ["涢泰饲料"],
        "zhongyan-dongxing": ["中盐东兴云梦制盐"], "guangyan-lantian": ["广盐蓝天"],
        "yunhong-pharma": ["云虹制药"], "luodong-bio": ["罗东生物"], "weidi-bio": ["威敌生物"],
    }

    def __init__(self, path: Path, base: dict[str, Any], year: int = 2025):
        super().__init__(path, base, year)
        self.original = deepcopy(base)
        self.data["regions"] = []
        self.data["chains"] = [dict(id=c["id"], name=c["name"], shortName=c["shortName"], stages=[]) for c in base["chains"]]
        self.data["profile"] = dict(name="云梦县", indicators=[], pillars=[], chainLinks=[], leaders=[], policies=[], parks=[])
        self.chain_lookup.update({norm("现代化工"): "chemicals", norm("现代化工/盐化工"): "chemicals", norm("新材料"): "new-energy"})
        self.coverage: dict[str, dict[str, list[int]]] = {}
        self.consumed_cells: set[tuple[str, str]] = set()
        self.company_rows: dict[str, list[int]] = {}
        self.raw_rows: dict[str, list[dict[str, Any]]] = {}
        self.merged_anchors: dict[tuple[str, str], Any] = {}
        for sheet in self.book:
            self.coverage[sheet.title] = {"headers": [], "data": [], "context": []}
            for merged in sheet.merged_cells.ranges:
                anchor = sheet.cell(merged.min_row, merged.min_col)
                for row in sheet.iter_rows(min_row=merged.min_row, max_row=merged.max_row, min_col=merged.min_col, max_col=merged.max_col):
                    for cell in row: self.merged_anchors[(sheet.title, cell.coordinate)] = anchor
            self.raw_rows[sheet.title] = [{"row": row[0].row, "cells": {cell.column_letter: {"value": plain(cell.value), "numberFormat": cell.number_format} for cell in row if cell.value is not None}} for row in sheet.iter_rows() if any(c.value is not None for c in row)]
        # The local records have no workbook coordinates. Keep logos and verified
        # legal detail, but discard their previous demonstration locations.
        for company in self.data["companies"]:
            if company["local"]:
                company.pop("coordinate", None)
                company["locationStatus"] = "missing"
                company["chainIds"] = []
                company["stageIds"] = []
                company["products"] = []

    def value(self, cell: Any) -> Any:
        if cell is not None and cell.value is None:
            cell = self.merged_anchors.get((cell.parent.title, cell.coordinate), cell)
        return super().value(cell)

    def mark(self, sheet: str, row: int, kind: str = "data"):
        if row not in self.coverage[sheet][kind]: self.coverage[sheet][kind].append(row)

    def field(self, row: tuple[Any, ...], header: dict[str, Any], *names: str):
        cell = super().field(row, header, *names)
        if cell is not None: self.consumed_cells.add((cell.parent.title, cell.coordinate))
        return cell

    def table(self, sheet: Any, required: list[tuple[str, ...]]):
        if sheet is None: return
        header = None
        for row in sheet.iter_rows():
            candidate = self.fields(row)
            if all(any(norm(name) in candidate for name in group) for group in required):
                header = candidate; self.mark(sheet.title, row[0].row, "headers"); continue
            if header is not None and any(not missing(c.value) for c in row):
                self.mark(sheet.title, row[0].row); yield row, header
        if header is None: self.errors.append({"sheet": sheet.title, "cell": "", "message": "缺少必需表头", "value": ""})

    def import_chains(self):
        sheet = self.sheet("统一产业链名单", ["产业链名单", "统一产业链"])
        for row, header in self.table(sheet, [("产业链",), ("环节（上/中/下游）", "环节"), ("包含哪些产品或业务",)]):
            get = lambda *names: self.field(row, header, *names)
            cid = self.chain(get("产业链"))
            position = self.text(get("环节（上/中/下游）", "环节"))
            if not cid: continue
            if position not in self.POSITION_IDS:
                self.issue(get("环节（上/中/下游）", "环节"), "环节必须是上游、中游或下游"); continue
            products = split_list(self.value(get("包含哪些产品或业务")))
            aliases = split_list(self.value(get("各地纲要里的常见说法")))
            sid = cid + "-" + self.POSITION_IDS[position]
            target = next(c for c in self.data["chains"] if c["id"] == cid)
            if any(s["id"] == sid for s in target["stages"]):
                self.issue(get("产业链"), "同一产业链位置重复"); continue
            name = "、".join(products[:2]) + ("等" if len(products) > 2 else "")
            target["stages"].append(dict(id=sid, stage=position, name=name or position, products=products, aliases=aliases, source=f"{self.path.name} · {sheet.title}第{row[0].row}行；显示名称为产品原文摘要"))
            for alias in aliases:
                # Parenthetical location lists are provenance, not chain names.
                self.chain_lookup.setdefault(norm(re.split(r"[（(]", alias)[0]), cid)
            self.counts["chainStages"] += 1

    @staticmethod
    def source_links(source: str, row: int) -> list[dict[str, str]]:
        urls = list(dict.fromkeys(re.findall(r"https?://[^\s｜|；;，、<>]+", source)))
        links = [dict(label=f"采集表 · 云梦画像第{row}行：{source.split('｜')[0].split('http')[0].rstrip('；;')}")]
        if urls: links[0]["url"] = urls[0]
        links.extend(dict(label=f"该行补充来源 {i + 2}", url=url) for i, url in enumerate(urls[1:]))
        return links

    def collected_companies(self, raw: str, chains: list[str], source: str, row: int, stage_id: str | None = None) -> list[str]:
        ids = []
        for cid, aliases in self.COMPANY_NAMES.items():
            matched = next((name for name in aliases if name in raw), None)
            if not matched: continue
            company = next((c for c in self.data["companies"] if c["id"] == cid), None)
            if company is None:
                company = dict(id=cid, name=aliases[-1], fullName=aliases[0], cityId="yunmeng", cityName="云梦县", logo=None, sector="待补充", summary="采集表列示的云梦企业或项目；具体法人名称、工商及坐标待核实。", chainIds=[], stageIds=[], products=[], local=True, relationship="existing", shareholders=[], sources=[], status="verified", locationStatus="missing")
                self.data["companies"].append(company); self.counts["newCompanies"] += 1
            if cid == "huasheng-xianghe":
                company["name"] = "华盛锂电（云梦项目）"
                company["summary"] = "采集表列示华盛锂电云梦项目；对应用户既有资本流中的湖北华盛祥和新能源材料有限公司，属于已有投资，非新未落地候选。"
            elif cid == "zhongyi":
                company["summary"] = "采集表及所列政府报告披露的云梦铜箔企业；工商、园区逐企业归属和精确位置待补充。"
            for chain_id in chains:
                if chain_id not in company["chainIds"]: company["chainIds"].append(chain_id)
            if stage_id and stage_id not in company["stageIds"]: company["stageIds"].append(stage_id)
            company["sector"] = "、".join(next(c["shortName"] for c in self.data["chains"] if c["id"] == chain_id) for chain_id in company["chainIds"]) or "产业归属待核实"
            # Product annotations belong to the named enterprise; generic sector
            # products must not be copied into each enterprise as business facts.
            annotation = re.search(re.escape(matched) + r"[（(]([^）)]+)[）)]", raw)
            if annotation and not any(x in annotation.group(1) for x in ("县政府报告", "具体企业")):
                company["products"] = list(dict.fromkeys(company["products"] + split_list(annotation.group(1))))
            if raw.startswith(matched + "："):
                company["products"] = list(dict.fromkeys(company["products"] + split_list(raw.split("：", 1)[1])))
            for link in self.source_links(source, row):
                if link not in company["sources"]: company["sources"].append(link)
            company["status"] = "verified"
            self.company_rows.setdefault(cid, []).append(row)
            ids.append(cid)
        return ids

    def import_profile(self):
        sheet = self.sheet("云梦画像")
        if sheet is None: return
        block = None; header = None
        for row in sheet.iter_rows():
            row_number = row[0].row
            if not any(c.value is not None for c in row): continue
            candidate = self.fields(row); has = lambda *names: any(norm(n) in candidate for n in names)
            kind = None
            if has("年份") and has("GDP（亿元）"): kind = "indicators"
            elif has("支柱产业") and has("产值规模（注明年度/口径）", "产值规模"): kind = "pillars"
            elif has("产业链（统一名称）", "产业链") and has("强链/补链"): kind = "chainLinks"
            elif has("产业链") and has("链长（姓名和职务）"): kind = "leaders"
            elif has("文件名") and has("发布单位"): kind = "policies"
            if kind: block, header = kind, candidate; self.mark(sheet.title, row_number, "headers"); continue
            if block is None: self.mark(sheet.title, row_number, "context"); continue
            get = lambda *names: self.field(row, header, *names)
            source = self.text(get("来源", "说明及来源", "公开信息及来源"), source=True)
            item = None
            if block == "indicators":
                year = self.number(get("年份"), integer=True, minimum=1900, maximum=2200)
                if year is None: continue
                item = dict(year=year, gdp=self.number(get("GDP（亿元）"), minimum=0), growth=self.number(get("GDP增速", "GDP增速（%）"), percentage=True), shares=[self.number(get(k, k+"（%）"), percentage=True, minimum=0, maximum=100) for k in ("一产占比", "二产占比", "三产占比")], industrialEnterprises=self.number(get("规上工业企业数（家）", "规上工业企业数"), minimum=0, integer=True), source=source, status="verified")
            elif block == "pillars":
                name = self.text(get("支柱产业")); cid = self.chain(get("支柱产业"))
                if not name or not cid: continue
                chain_ids = ["new-energy", "chemicals"] if name == "新材料" else [cid]
                raw_companies = self.text(get("龙头企业/重点企业", "龙头企业"))
                # A broad pillar can span two chains without placing every
                # listed enterprise in both. Specific county-stage evidence
                # below supplies the enterprise memberships for 新材料.
                company_chain_ids = [] if name == "新材料" else chain_ids
                item = dict(id="yunmeng-"+cid, name=name, chainId=cid, chainIds=chain_ids, output=self.text(get("产值规模（注明年度/口径）", "产值规模")) or None, companyIds=self.collected_companies(raw_companies, company_chain_ids, source, row_number), park=self.text(get("所在园区/载体", "所在园区")) or None, source=source)
            elif block == "chainLinks":
                cid = self.chain(get("产业链（统一名称）", "产业链")); position = self.text(get("环节"))
                if not cid: continue
                if position not in self.POSITION_IDS: self.issue(get("环节"), "画像环节需要上游、中游或下游"); continue
                sid = cid + "-" + self.POSITION_IDS[position]
                reported = self.text(get("强链/补链")); basis = plain(self.value(get("云梦企业/现有基础", "企业名")))
                if reported not in ("强链", "补链"): self.issue(get("强链/补链"), "需明确原表强链/补链标签"); continue
                pending = reported == "补链" and bool(re.search(r"未找到|未核实|待核|缺少.*名录", source))
                item = dict(stageId=sid, status="unknown" if pending else "present" if reported == "强链" else "gap", reportedStatus=reported, evidenceStatus="pending" if pending else "reported", basis=basis, companyIds=self.collected_companies(basis, [cid], source, row_number, sid), source=source)
            elif block == "leaders":
                cid = self.chain(get("产业链"))
                if not cid: continue
                raw_person = plain(self.value(get("链长（姓名和职务）"))); raw_company = plain(self.value(get("链主企业")))
                notes = source + f"；原表链长备注：{raw_person}；原表链主备注：{raw_company}"
                item = dict(chainId=cid, name=None if missing(raw_person) else raw_person, title=None, company=None if missing(raw_company) else raw_company, source=notes)
            elif block == "policies":
                name = self.text(get("文件名")); raw_link = self.text(get("链接/用途", "链接"), source=True)
                links = re.findall(r"https?://[^\s｜|；;，、<>]+", raw_link)
                if not name: continue
                item = dict(id=stable_id("policy", name), name=name, publisher=self.text(get("发布单位")), date=self.text(get("发布日期")), url=links[0] if links else None, source=source or raw_link or name)
            if item is not None:
                self.data["profile"][block].append(item); self.counts[block] += 1; self.mark(sheet.title, row_number)
        park_candidates = [
            dict(id="yunmeng-salt-chemical-park", name="云梦盐化工产业园", source="采集表云梦画像第13–14行及所列经信厅来源", status="verified", planning=False, note="具体企业归属未逐家核实，未提供园区边界或坐标"),
            dict(id="yunmeng-development-zone", name="云梦经济开发区", source="采集表云梦画像第14–15行", status="verified", planning=False, note="原表说明企业园区归属未逐企业核实；未提供精确范围"),
            dict(id="yunmeng-agri-processing-plan", name="万亩百亿农产品加工示范区", source="采集表云梦画像第16行：县政府工作报告", status="missing", planning=True, note="规划目标，非已建成园区；具体园区缺"),
        ]
        carrier_text = "\n".join(p.get("park") or "" for p in self.data["profile"]["pillars"])
        self.data["profile"]["parks"] = [p for p in park_candidates if p["name"] in carrier_text]

    def run(self) -> tuple[dict[str, Any], dict[str, Any]]:
        self.import_chains(); self.import_regions(); self.import_profile()
        # Independent, already sourced company leads remain outside the workbook
        # claim. Bind them to the actual catalog products, not deleted demo IDs.
        valid_stages = {s["id"] for c in self.data["chains"] for s in c["stages"]}
        for company in self.data["companies"]:
            supplemental = ["new-energy-midstream", "chemicals-downstream"] if company["id"] == "capchem" else ["new-energy-midstream"] if company["id"] in ("tinci", "huasheng") else []
            if supplemental:
                company["stageIds"] = list(dict.fromkeys([s for s in company["stageIds"] if s in valid_stages] + [s for s in supplemental if s in valid_stages]))
        province_count = sum(r["level"] == "省" for r in self.data["regions"])
        city_count = sum(r["level"] == "市" for r in self.data["regions"])
        pending_count = sum(x.get("evidenceStatus") == "pending" for x in self.data["profile"]["chainLinks"])
        self.data["meta"].update(updatedAt=datetime.now(timezone.utc).isoformat(timespec="seconds"), year=self.year, description=f"全国指标、原文产业及云梦画像全面接入用户提供的《{self.path.name}》：{province_count}省+{city_count}市、{len(self.data['chains'])}产业链{len(valid_stages)}个环节。数值保留原文来源、缺失为空；采集来源未作独立审计。{pending_count}项原表补链仅据未找到/未核实资料，保留原标签并标待核实。城市匹配分数仍为模型示意，真实流线资源未变。新宙邦/天赐等独立公开企业线索保留官网来源，不声称由本表提供。")
        self.data["profile"]["indicators"].sort(key=lambda x: x["year"])
        summaries = []
        for sheet in self.book:
            nonempty = {row[0].row for row in sheet.iter_rows() if any(c.value is not None for c in row)}
            seen = set(sum(self.coverage[sheet.title].values(), []))
            unhandled = sorted(nonempty - seen)
            if unhandled: self.errors.append({"sheet": sheet.title, "cell": "", "message": f"存在未处理的非空行 {unhandled}", "value": ""})
            unhandled_cells = [cell.coordinate for row in sheet.iter_rows() if row[0].row in self.coverage[sheet.title]["data"] for cell in row if cell.value is not None and (sheet.title, cell.coordinate) not in self.consumed_cells]
            if unhandled_cells: self.errors.append({"sheet": sheet.title, "cell": ",".join(unhandled_cells), "message": "数据行存在未处理的非空单元格", "value": ""})
            summaries.append(dict(sheet=sheet.title, rows=sheet.max_row, columns=sheet.max_column, mergedRanges=[str(x) for x in sheet.merged_cells.ranges], nonemptyRows=len(nonempty), importedRows=self.coverage[sheet.title]["data"], headerRows=self.coverage[sheet.title]["headers"], contextRows=self.coverage[sheet.title]["context"], unhandledRows=unhandled, unhandledColumns=sorted(set(re.sub(r"\d", "", cell) for cell in unhandled_cells)), unhandledCells=unhandled_cells))
        numeric = ["gdp", "growth"]
        missing_counts = {key: sum(r[key] is None for r in self.data["regions"]) for key in numeric}
        missing_counts["anyShare"] = sum(any(v is None for v in r["shares"]) for r in self.data["regions"])
        missing_counts["pillars"] = sum(not r["pillars"] for r in self.data["regions"])
        missing_counts["priorities"] = sum(not r["priorities"] for r in self.data["regions"])
        report = dict(input=str(self.path.resolve()), mode="replace-collected", counts=self.counts, errors=self.errors, warnings=self.warnings, outputWritten=False, sheets=summaries, missingRegionFields=missing_counts, chainEvidence={"present": sum(x["status"] == "present" for x in self.data["profile"]["chainLinks"]), "pendingReportedGaps": sum(x.get("evidenceStatus") == "pending" for x in self.data["profile"]["chainLinks"])}, companySourceRows=self.company_rows, independentCompanies=[dict(id=c["id"], name=c["name"], sources=c["sources"]) for c in self.data["companies"] if c["id"] not in self.company_rows], rawWorkbook=self.raw_rows, notes=["统一产业链合并单元格按锚点值读取，未修改原 Excel。", "产业环节显示名是产品列表摘要，完整原文保留在 products/aliases。", "资源、基地、平台与项目说明保留在 basis；未将其全部伪造为法人企业。", "未提供注册资本、法人、股东、logo、精确工厂坐标；已有已核对字段及logo保持，其他缺失。", "实际产值、同比增速及2030目标完整保留在产值规模原文，不把目标值当实际产值。"])
        self.book.close(); self.cached.close()
        return self.data, report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="一个 Excel，包含约定的三个 sheet")
    parser.add_argument("--output", required=True, type=Path, help="输出 JSON，建议先输出到 work/ 审阅")
    parser.add_argument("--base", type=Path, default=DEFAULT_BASE, help="合并基线；企业、城市分数、logo 保留")
    parser.add_argument("--report", type=Path, help="报告路径；默认输出文件旁 .report.json")
    parser.add_argument("--year", type=int, default=2025)
    parser.add_argument("--replace", action="store_true", help="使用实际采集表全面替换省市、产业环节和县画像；保留独立企业线索及模型评分")
    args = parser.parse_args(argv)
    report_path = args.report or args.output.with_suffix(".report.json")
    try:
        base = json.loads(args.base.read_text(encoding="utf-8-sig"))
        importer = CollectedWorkbookImporter if args.replace else Importer
        data, report = importer(args.input, base, args.year).run()
    except Exception as exc:
        report = {"input": str(args.input), "errors": [{"message": str(exc)}], "warnings": [], "outputWritten": False}
        data = None
    if not report["errors"] and data is not None:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        temp = args.output.with_name(args.output.name + ".tmp")
        temp.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        temp.replace(args.output)
        report["outputWritten"] = True
        report["output"] = str(args.output.resolve())
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"导入{'完成' if report['outputWritten'] else '未写入'}：{len(report['errors'])} 个错误，{len(report['warnings'])} 个提示。报告：{report_path}")
    return 0 if report["outputWritten"] else 2


if __name__ == "__main__":
    sys.exit(main())
