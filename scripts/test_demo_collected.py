"""Compare the delivered workbook with the actual runtime dataset, cell by cell."""
import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from import_demo_excel import CollectedWorkbookImporter, missing, split_list

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--input", required=True, type=Path)
parser.add_argument("--data", type=Path, default=Path("public/assets/demo/demo-data.json"))
parser.add_argument("--report", type=Path, default=Path("public/assets/demo/data-import-report.json"))
args = parser.parse_args()
book = load_workbook(args.input, data_only=True)
data = json.loads(args.data.read_text(encoding="utf-8"))
report = json.loads(args.report.read_text(encoding="utf-8"))
checks = 0


def equal(actual, expected, message):
    global checks
    assert actual == expected, f"{message}: {actual!r} != {expected!r}"
    checks += 1


def number(cell, percent=False):
    if missing(cell.value): return None
    value = float(cell.value)
    if percent and "%" in cell.number_format: value *= 100
    return round(value, 10)


equal(len(data["regions"]), 131, "regions")
equal(sum(r["level"] == "省" for r in data["regions"]), 31, "provinces")
equal(sum(r["level"] == "市" for r in data["regions"]), 100, "cities")
for row in book["全国省市"].iter_rows(min_row=2):
    record = next(r for r in data["regions"] if r["name"] == row[0].value and r["level"] == row[1].value)
    for key, index, percent in [("gdp", 3, False), ("growth", 4, True)]: equal(record[key], number(row[index], percent), f"row {row[0].row} {key}")
    equal(record["shares"], [number(c, True) for c in row[5:8]], f"row {row[0].row} shares")
    equal(record["pillars"], split_list(row[8].value), f"row {row[0].row} pillars")
    equal(record["priorities"], split_list(row[9].value), f"row {row[0].row} priorities")
    equal(record["source"], row[10].value, f"row {row[0].row} source")
equal(sum(len(c["stages"]) for c in data["chains"]), 30, "stages")
last_name = None
for row in book["统一产业链"].iter_rows(min_row=2):
    if row[0].value: last_name = row[0].value
    chain = next(c for c in data["chains"] if c["name"] == last_name)
    stage = next(s for s in chain["stages"] if s["stage"] == row[1].value)
    equal(stage["products"], split_list(row[2].value), f"chain row {row[0].row} products")
    equal(stage["aliases"], split_list(row[3].value), f"chain row {row[0].row} aliases")
sheet = book["云梦画像"]
for row in sheet.iter_rows(min_row=5, max_row=9):
    record = next(r for r in data["profile"]["indicators"] if r["year"] == row[0].value)
    equal(record["gdp"], number(row[1]), f"county {row[0].value} gdp")
    equal(record["growth"], number(row[2], True), f"county {row[0].value} growth")
    equal(record["shares"], [number(c, True) for c in row[3:6]], f"county {row[0].value} shares")
    equal(record["industrialEnterprises"], row[6].value, f"county {row[0].value} industry")
    equal(record["source"], row[7].value, f"county {row[0].value} source")
equal([p["name"] for p in data["profile"]["pillars"]], [sheet.cell(r, 1).value for r in range(13, 17)], "pillars")
for record, row in zip(data["profile"]["pillars"], range(13, 17)):
    equal(record["output"], sheet.cell(row, 2).value, f"pillar {row} actual/target wording")
    equal(record["park"], sheet.cell(row, 4).value, f"pillar {row} park qualifications")
equal(len(data["profile"]["chainLinks"]), 18, "county links")
equal(sum(x["status"] == "present" for x in data["profile"]["chainLinks"]), 10, "reported strengths")
equal(sum(x.get("evidenceStatus") == "pending" for x in data["profile"]["chainLinks"]), 8, "pending gaps")
for link, row in zip(data["profile"]["chainLinks"], range(21, 39)):
    equal(link["reportedStatus"], sheet.cell(row, 3).value, f"link {row} original status")
    equal(link["basis"], sheet.cell(row, 4).value, f"link {row} basis")
    equal(link["source"], sheet.cell(row, 5).value, f"link {row} source")
    if link["reportedStatus"] == "补链": equal(link["status"], "unknown", f"link {row} not-found is not absence")
equal(len(data["profile"]["leaders"]), 5, "leaders")
equal(all(x["name"] is None and x["company"] is None for x in data["profile"]["leaders"]), True, "unconfirmed leaders remain missing")
equal(len(data["profile"]["policies"]), 5, "policies")
for policy, row in zip(data["profile"]["policies"], range(51, 56)):
    equal(policy["name"], sheet.cell(row, 1).value, f"policy {row} name")
    equal(policy["date"], sheet.cell(row, 3).value.strftime("%Y-%m-%d"), f"policy {row} date")
    equal(policy["source"], sheet.cell(row, 4).value, f"policy {row} scope")
equal(len([c for c in data["companies"] if c["local"]]), 21, "named local companies/projects")
equal(all("coordinate" not in c for c in data["companies"] if c["local"]), True, "no invented local coordinates")
equal(any(s["id"] == "electrolyte" for c in data["chains"] for s in c["stages"]), False, "deleted example stage")
equal(data["cities"][0]["scores"], [94, 90, 92, 92], "independent sample scores preserved")
companies = {c["id"]: c for c in data["companies"]}
equal(companies["zhongyi"]["chainIds"], ["new-energy", "electronics"], "broad materials pillar does not make copper foil a chemicals company")
equal(companies["huaxin-silicone"]["chainIds"], ["chemicals"], "broad materials pillar does not prove an energy chain membership")
equal(next(p for p in data["profile"]["pillars"] if p["name"] == "新材料")["chainIds"], ["new-energy", "chemicals"], "pillar retains cross-chain scope")
# Exercise the actual full replacement path, including the supplemental website
# evidence and image resources added after the initial workbook import.
reimported, reimport_report = CollectedWorkbookImporter(args.input, data).run()
equal(reimport_report["errors"], [], "replacement regression")
reimported_companies = {c["id"]: c for c in reimported["companies"]}
equal(reimported_companies["zhongyi"]["chainIds"], ["new-energy", "electronics"], "replacement keeps specific copper foil membership")
equal(reimported_companies["huaxin-silicone"]["chainIds"], ["chemicals"], "replacement avoids broad-pillar propagation")
equal(reimported_companies["capchem"], companies["capchem"], "supplemental Capchem stages, sources and logo preserved")
for cid, company in companies.items():
    equal(reimported_companies[cid]["logo"], company["logo"], f"replacement logo preservation {cid}")
equal(all(not s["unhandledRows"] and not s["unhandledColumns"] and not s["unhandledCells"] for s in report["sheets"]), True, "coverage")
if "sourceWorkbookSha256" in report: equal(hashlib.sha256(args.input.read_bytes()).hexdigest(), report["sourceWorkbookSha256"], "source identity")
book.close()
print(f"Collected workbook: {checks} source/coverage checks passed.")
