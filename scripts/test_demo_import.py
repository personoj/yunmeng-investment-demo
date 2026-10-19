"""Focused regression tests for Excel ingestion; uses disposable workbook copies."""
import contextlib
import io
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from import_demo_excel import DEFAULT_BASE, Importer, main, split_list

TEMPLATE = DEFAULT_BASE.parent / "数据模板.xlsx"


class DemoImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "collection.xlsx"
        self.base = json.loads(DEFAULT_BASE.read_text(encoding="utf-8"))

    def workbook(self, changes=None):
        book = load_workbook(TEMPLATE)
        if changes: changes(book)
        book.save(self.path); book.close()
        return self.path

    def run_import(self, changes=None):
        self.workbook(changes)
        return Importer(self.path, self.base).run()

    def test_blank_template_roundtrip_preserves_independent_assets(self):
        data, report = self.run_import()
        self.assertEqual(report["errors"], [])
        self.assertEqual(report["counts"]["regions"], 31)
        self.assertEqual(report["counts"]["chainStages"], sum(len(c["stages"]) for c in self.base["chains"]))
        self.assertEqual(data["companies"], self.base["companies"])
        self.assertEqual(data["cities"], self.base["cities"])
        self.assertIsNone(next(r for r in data["regions"] if r["id"] == "hubei")["gdp"])
        expected = next(s for c in self.base["chains"] if c["id"] == "new-energy" for s in c["stages"] if s["stage"] == "中游")
        self.assertEqual(next(s for c in data["chains"] if c["id"] == "new-energy" for s in c["stages"] if s["id"] == expected["id"])["name"], expected["name"])

    def test_zero_missing_percentage_and_source_fidelity(self):
        def change(book):
            sheet = book["全国省市数据表"]
            sheet["D6"] = 0; sheet["E6"] = 0.058; sheet["E6"].number_format = "0.0%"
            sheet["F6"] = "8.5%"; sheet["G6"] = "缺"; sheet["H6"] = "—"
            sheet["I6"] = "氯碱化工、光电子信息"; sheet["J6"] = "未来产业（低空经济、人工智能等）"
            sheet["K6"] = "湖北省 2025 年统计公报 第 3 页"; sheet["K6"].hyperlink = "https://example.org/source"
            sheet["L6"] = "化工新材料（含盐化工）、电子信息与半导体、未来产业（低空经济、人工智能等）"
        data, report = self.run_import(change)
        self.assertEqual(report["errors"], [])
        row = next(r for r in data["regions"] if r["id"] == "hubei")
        self.assertEqual(row["gdp"], 0)
        self.assertEqual(row["growth"], 5.8)
        self.assertEqual(row["shares"], [8.5, None, None])
        self.assertEqual(row["pillars"], ["氯碱化工", "光电子信息"])
        self.assertEqual(row["chainIds"], ["chemicals", "electronics", "future"])
        self.assertEqual(row["source"], "湖北省 2025 年统计公报 第 3 页 | https://example.org/source")

    def test_malformed_numeric_does_not_overwrite_output(self):
        self.workbook(lambda b: setattr(b["全国省市数据表"]["D6"], "value", "约六万亿元"))
        output = Path(self.temp.name) / "existing.json"
        output.write_text("keep-existing", encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()):
            code = main(["--input", str(self.path), "--output", str(output)])
        self.assertEqual(code, 2)
        self.assertEqual(output.read_text(encoding="utf-8"), "keep-existing")
        report = json.loads(output.with_suffix(".report.json").read_text(encoding="utf-8"))
        self.assertFalse(report["outputWritten"])
        self.assertTrue(any("数字格式错误" in x["message"] for x in report["errors"]))

    def test_unknown_explicit_mapping_is_error_raw_alias_is_warning(self):
        _, report = self.run_import(lambda b: setattr(b["全国省市数据表"]["L6"], "value", "未定义产业链"))
        self.assertTrue(any("未映射" in x["message"] for x in report["errors"]))
        data, report = self.run_import(lambda b: setattr(b["全国省市数据表"]["I6"], "value", "本地特色产业原话"))
        self.assertEqual(report["errors"], [])
        self.assertTrue(any("本地特色产业原话" in x["message"] for x in report["warnings"]))
        self.assertIn("本地特色产业原话", next(r for r in data["regions"] if r["id"] == "hubei")["pillars"])

    def test_profile_five_blocks_and_new_company_keep_logo(self):
        self.base["companies"][0]["logo"] = "assets/demo/company-logos/custom.png"
        def change(book):
            sheet = book["云梦画像"]
            sheet["B11"] = 0; sheet["G11"] = 0; sheet["H11"] = "统计公报"
            for col, value in zip("ABCDEF", ["新能源及材料", "新能源及材料", "36 亿元", "中一科技、测试新增企业", "云梦县经济开发区", "工作报告"]): sheet[f"{col}16"] = value
            sheet["C24"] = "强链"; sheet["D24"] = "中一科技"; sheet["E24"] = "工作报告"
            sheet["B45"] = "待补充"; sheet["C45"] = "分管职务待补充"; sheet["E45"] = "链长制文件"
            for col, value in zip("ABCDE", ["示例政策文件", "发布单位", "2025-01-01", "https://example.org/policy", "政府网站"]): sheet[f"{col}58"] = value
        data, report = self.run_import(change)
        self.assertEqual(report["errors"], [])
        self.assertEqual(report["counts"]["newCompanies"], 1)
        self.assertEqual(data["companies"][0]["logo"], "assets/demo/company-logos/custom.png")
        added = next(c for c in data["companies"] if c["name"] == "测试新增企业")
        self.assertEqual(added["shareholders"], []); self.assertIsNone(added["logo"])
        self.assertEqual(data["profile"]["indicators"][-1]["industrialEnterprises"], 0)
        self.assertEqual(next(p for p in data["profile"]["policies"] if p["name"] == "示例政策文件")["url"], "https://example.org/policy")
        stage_id = next(s["id"] for c in self.base["chains"] if c["id"] == "new-energy" for s in c["stages"] if s["stage"] == "中游")
        self.assertEqual(next(x for x in data["profile"]["chainLinks"] if x["stageId"] == stage_id)["companyIds"], ["zhongyi"])

    def test_duplicate_regions_and_uncached_formulas_are_errors(self):
        def change(book):
            s = book["全国省市数据表"]
            for col in range(1, 13): s.cell(7, col).value = s.cell(6, col).value
            s["D6"] = "=100+200"
        _, report = self.run_import(change)
        self.assertTrue(any("重复" in x["message"] for x in report["errors"]))
        self.assertTrue(any("缓存" in x["message"] for x in report["errors"]))

    def test_parenthesis_lists_remain_one_chain(self):
        self.assertEqual(split_list("新能源及材料、未来产业（低空经济、人工智能等）"), ["新能源及材料", "未来产业（低空经济、人工智能等）"])


if __name__ == "__main__":
    unittest.main()
